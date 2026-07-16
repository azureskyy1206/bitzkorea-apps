"""관리자 SDL CRUD·엑셀 적재."""

import io
import uuid

import pytest
from httpx import AsyncClient
from openpyxl import Workbook


@pytest.mark.asyncio
async def test_admin_sdl_crud(client: AsyncClient, admin_headers: dict[str, str]):
    suffix = uuid.uuid4().hex[:8]
    create_school = await client.post(
        "/admin/schools",
        headers=admin_headers,
        json={"name": f"SDL테스트학교{suffix}", "biz_no": f"333-33-{suffix[:5]}"},
    )
    assert create_school.status_code == 201
    school_id = create_school.json()["id"]

    listed = await client.get("/admin/sdl", headers=admin_headers)
    assert listed.status_code == 200

    created = await client.post(
        "/admin/sdl",
        headers=admin_headers,
        json={
            "school_id": school_id,
            "registered_ip": "192.168.1.10",
            "purchase_qty": 25,
            "contact_name": "테스트담당",
            "contact_phone": "010-1234-5678",
            "contact_email": f"sdl-{suffix}@school.kr",
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["school_name"].startswith("SDL테스트학교")
    assert body["registered_ip"] == "192.168.1.10"
    assert body["purchase_qty"] == 25
    sdl_id = body["id"]

    listed2 = await client.get(
        "/admin/sdl",
        params={"q": "192.168"},
        headers=admin_headers,
    )
    assert listed2.status_code == 200
    assert listed2.json()["total"] == 1

    patched = await client.patch(
        f"/admin/sdl/{sdl_id}",
        headers=admin_headers,
        json={"purchase_qty": 30, "registered_ip": "192.168.1.11"},
    )
    assert patched.status_code == 200
    assert patched.json()["purchase_qty"] == 30
    assert patched.json()["registered_ip"] == "192.168.1.11"

    deleted = await client.delete(f"/admin/sdl/{sdl_id}", headers=admin_headers)
    assert deleted.status_code == 204


@pytest.mark.asyncio
async def test_admin_sdl_excel_import(client: AsyncClient, admin_headers: dict[str, str]):
    suffix = uuid.uuid4().hex[:8]
    wb = Workbook()
    ws = wb.active
    ws.append(["학교명", "담당자", "연락처", "이메일", "등록 IP", "SDL 구매수량"])
    ws.append(
        [
            f"SDL엑셀학교{suffix}",
            "엑셀담당",
            "010-9999-8888",
            f"excel-{suffix}@school.kr",
            "10.0.0.50",
            "15",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = await client.post(
        "/admin/import/sdl",
        headers=admin_headers,
        files={"file": ("sdl_test.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"replace_all": "false"},
    )
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["created"] >= 1

    listed = await client.get("/admin/sdl", params={"q": f"SDL엑셀학교{suffix}"}, headers=admin_headers)
    assert listed.status_code == 200
    assert listed.json()["total"] >= 1
    item = listed.json()["items"][0]
    assert item["registered_ip"] == "10.0.0.50"
    assert item["purchase_qty"] == 15


@pytest.mark.asyncio
async def test_admin_sdl_grouped_sort_and_export(client: AsyncClient, admin_headers: dict[str, str]):
    suffix = uuid.uuid4().hex[:8]
    create_a = await client.post(
        "/admin/schools",
        headers=admin_headers,
        json={"name": f"AAA SDL{suffix}", "biz_no": f"111-11-{suffix[:5]}"},
    )
    create_b = await client.post(
        "/admin/schools",
        headers=admin_headers,
        json={"name": f"ZZZ SDL{suffix}", "biz_no": f"222-22-{suffix[:5]}"},
    )
    school_a = create_a.json()["id"]
    school_b = create_b.json()["id"]

    for school_id, ip, qty in [
        (school_a, "10.1.0.1", 10),
        (school_b, "10.2.0.1", 50),
    ]:
        res = await client.post(
            "/admin/sdl",
            headers=admin_headers,
            json={"school_id": school_id, "registered_ip": ip, "purchase_qty": qty},
        )
        assert res.status_code == 201, res.text

    # 동일 학교·담당자에 IP 2개 → 엑셀 1행, IP 1·IP 2 열
    multi = await client.post(
        "/admin/sdl",
        headers=admin_headers,
        json={
            "school_id": school_a,
            "registered_ip": "10.1.0.2",
            "purchase_qty": 0,
            "contact_name": "멀티IP",
            "contact_email": f"multi-{suffix}@school.kr",
        },
    )
    assert multi.status_code == 201, multi.text
    await client.post(
        "/admin/sdl",
        headers=admin_headers,
        json={
            "school_id": school_a,
            "registered_ip": "10.1.0.3",
            "purchase_qty": 10,
            "contact_name": "멀티IP",
            "contact_email": f"multi-{suffix}@school.kr",
        },
    )

    by_qty = await client.get(
        "/admin/sdl/grouped",
        params={"q": suffix, "sort_by": "total_qty", "sort_order": "desc"},
        headers=admin_headers,
    )
    assert by_qty.status_code == 200
    items = by_qty.json()["items"]
    assert len(items) == 2
    assert items[0]["total_qty"] >= items[1]["total_qty"]

    export_res = await client.get(
        "/admin/sdl/export",
        params={"q": suffix},
        headers=admin_headers,
    )
    assert export_res.status_code == 200
    assert "spreadsheetml" in export_res.headers.get("content-type", "")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(export_res.content))
    ws = wb.active
    headers_row = [cell.value for cell in ws[1]]
    assert headers_row == [
        "학교명",
        "담당자",
        "연락처",
        "이메일",
        "IP 1",
        "IP 2",
        "IP 3",
        "IP 4",
        "IP 5",
        "SDL 구매수량",
    ]
    data_rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 3
    aaa_rows = [r for r in data_rows if str(r[0]).startswith("AAA SDL")]
    assert len(aaa_rows) == 2
    multi_row = [r for r in aaa_rows if r[1] == "멀티IP"]
    assert len(multi_row) == 1
    assert multi_row[0][4] == "10.1.0.2"
    assert multi_row[0][5] == "10.1.0.3"
    assert multi_row[0][9] == 10
    bare_row = [r for r in aaa_rows if not r[1]]
    assert len(bare_row) == 1
    assert bare_row[0][4] == "10.1.0.1"
