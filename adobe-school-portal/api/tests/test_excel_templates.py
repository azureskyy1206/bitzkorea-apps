"""엑셀 양식 템플릿 API."""

import io

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.services.excel_templates import (
    MASTER_LIST_COLUMNS,
    NAMED_MAIL_COLUMNS,
    build_excel_template_xlsx,
)


@pytest.mark.parametrize(
    "key,expected_columns",
    [
        ("master-list", MASTER_LIST_COLUMNS),
        ("named-mail", NAMED_MAIL_COLUMNS),
        (
            "usage-accounts",
            (
                "학교명",
                "사업자번호",
                "담당자",
                "담당자 연락처",
                "담당자 이메일",
                "라이선스 기간",
                "라이선스 할당량",
                "라이선스 할당/미할당",
                "할당 합계",
                "사용계정",
            ),
        ),
        (
            "sdl",
            (
                "학교명",
                "담당자",
                "연락처",
                "이메일",
                "IP 1",
                "IP 2",
                "IP 3",
                "IP 4",
                "IP 5",
                "SDL 구매수량",
            ),
        ),
    ],
)
def test_build_excel_template_xlsx(key: str, expected_columns: tuple[str, ...]) -> None:
    xlsx = build_excel_template_xlsx(key)
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == list(expected_columns)
    assert ws.max_row >= 2


@pytest.mark.asyncio
async def test_download_excel_template_api(client: AsyncClient, admin_headers: dict[str, str]) -> None:
    res = await client.get("/admin/excel-templates/named-mail", headers=admin_headers)
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(res.content))
    assert [cell.value for cell in wb.active[1]] == list(NAMED_MAIL_COLUMNS)


@pytest.mark.asyncio
async def test_download_excel_template_unknown(client: AsyncClient, admin_headers: dict[str, str]) -> None:
    res = await client.get("/admin/excel-templates/unknown", headers=admin_headers)
    assert res.status_code == 404
