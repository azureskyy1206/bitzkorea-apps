#!/usr/bin/env python3
import openpyxl
from pathlib import Path

KEYWORDS = ("화계", "김지원", "이행학", "210-83-00293")

def scan(path: Path, limit: int = 30) -> list:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    hits = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            s = " ".join(str(c) for c in row if c is not None)
            if any(k in s for k in KEYWORDS):
                hits.append((sn, i, row))
                if len(hits) >= limit:
                    wb.close()
                    return hits
    wb.close()
    return hits

def main() -> None:
    base = Path("/data/in")
    for name in sorted(base.glob("*.xlsx")):
        if "CY26" not in name.name and "cy26" not in name.name.lower():
            continue
        hits = scan(name)
        if not hits:
            continue
        print(f"\n=== {name.name} ({len(hits)} hits) ===")
        for sn, i, row in hits:
            print(f"  [{sn}] row {i}: {row[:18]}")

if __name__ == "__main__":
    main()
