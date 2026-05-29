from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side


NAVY = "1F4E78"
YELLOW = "FFFF00"
HEADERS = ["매출월", "담당자", "고객사", "제품명", "금액 변동 사유",
           "매출월 변동", "이익액 변동", "Loss or Drop"]
NCOL = len(HEADERS)


def build_excel(business_date: str, rows: list[dict], totals: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "일일 변경 내역"

    header_fill = PatternFill("solid", fgColor=NAVY)
    sum_fill = PatternFill("solid", fgColor=YELLOW)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.cell(1, 1, f"{business_date} (전일 대비)")
    ws.cell(1, 1).font = Font(bold=True, color="FFFFFF", size=11)
    ws.cell(1, 1).fill = header_fill
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)

    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(2, i, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for ri, row in enumerate(rows, start=3):
        profit = int(row.get("이익액_변동") or 0)
        loss = int(row.get("loss_drop") or 0)
        prod = row.get("제품명") or row.get("Program") or "-"
        vals = [
            row.get("매출월"), row.get("담당자"), row.get("고객사"),
            prod, row.get("사유"), row.get("매출월변동") or "-",
            profit if profit else "-", loss if loss else "-",
        ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v)
            c.border = border
            if ci in (7, 8) and isinstance(v, (int, float)):
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
            elif ci in (1, 2, 6):
                c.alignment = Alignment(horizontal="center")

    last = len(rows) + 2
    sum_r = last + 1
    ws.cell(sum_r, 1, "합계")
    ws.merge_cells(start_row=sum_r, start_column=1, end_row=sum_r, end_column=6)
    ws.cell(sum_r, 7, f"=SUM(G3:G{last})")
    ws.cell(sum_r, 8, f"=SUM(H3:H{last})")
    for i in range(1, NCOL + 1):
        c = ws.cell(sum_r, i)
        c.fill = sum_fill
        c.font = Font(bold=True)
        c.border = border
        if i in (7, 8):
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
        else:
            c.alignment = Alignment(horizontal="center")

    net_r = sum_r + 1
    ws.cell(net_r, 6, "순합계")
    ws.cell(net_r, 6).fill = sum_fill
    ws.cell(net_r, 6).font = Font(bold=True)
    ws.cell(net_r, 6).alignment = Alignment(horizontal="center")
    ws.cell(net_r, 6).border = border
    ws.cell(net_r, 7, f"=G{sum_r}-H{sum_r}")
    c = ws.cell(net_r, 7)
    c.number_format = "#,##0"; c.font = Font(bold=True); c.fill = sum_fill
    c.alignment = Alignment(horizontal="right"); c.border = border

    widths = [10, 10, 24, 40, 38, 16, 16, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    # 제품명 컬럼(D=4) 기본 숨김
    ws.column_dimensions["D"].hidden = True

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
