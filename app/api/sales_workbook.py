"""영업관리 xlsx — 현황판·2026 시트만 경량 로드/저장 (브라우저 전체 변환 회피)."""
from __future__ import annotations

import gzip
import json
import logging
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger(__name__)

# Fortune Sheet 0-indexed celldata 범위 (1-based 엑셀 좌표 기준 상한)
EDITOR_BOUNDS: dict[str, tuple[int, int, int, int]] = {
    "현황판": (1, 400, 1, 45),
    "2026": (1, 1300, 1, 35),
}

EDITOR_SHEET_NAMES = list(EDITOR_BOUNDS.keys())


def editor_cache_path(upload_dir: Path, batch_id: int, sheet_name: str) -> Path:
    d = upload_dir / ".editor-cache"
    d.mkdir(parents=True, exist_ok=True)
    safe = sheet_name.replace("/", "_")
    return d / f"batch_{batch_id}_{safe}.json.gz"


def _cache_is_fresh(cache: Path, xlsx: Path) -> bool:
    if not cache.is_file():
        return False
    try:
        return cache.stat().st_mtime >= xlsx.stat().st_mtime
    except OSError:
        return False


def load_cached_sheet(cache: Path) -> dict[str, Any] | None:
    if not cache.is_file():
        return None
    try:
        raw = gzip.decompress(cache.read_bytes())
        return json.loads(raw.decode("utf-8"))
    except Exception as e:
        logger.warning("editor cache read failed %s: %s", cache, e)
        return None


def write_cached_sheet(cache: Path, sheet: dict[str, Any]) -> None:
    cache.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(sheet, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    cache.write_bytes(gzip.compress(payload, compresslevel=6))


def _cell_display(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if v.time() else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return ""
        if v == int(v):
            return str(int(v))
    return str(v)


def _fortune_cell_from_openpyxl(cell: Any) -> dict[str, Any] | None:
    if cell.value is None and cell.data_type != "f":
        return None

    out: dict[str, Any] = {}
    if cell.data_type == "f":
        f = cell.value
        if f is None:
            return None
        fs = str(f).strip()
        if fs and not fs.startswith("="):
            fs = "=" + fs
        out["f"] = fs
    elif isinstance(cell.value, (int, float)):
        out["v"] = cell.value
        out["ct"] = {"fa": "General", "t": "n"}
    elif isinstance(cell.value, (datetime, date)):
        out["v"] = _cell_display(cell.value)
        out["ct"] = {"fa": "General", "t": "s"}
    else:
        out["v"] = cell.value
        out["ct"] = {"fa": "General", "t": "s"}

    out["m"] = _cell_display(cell.value) if cell.data_type != "f" else out.get("m", "")
    return out


def export_editor_sheet(path: Path, sheet_name: str, *, active: bool = False) -> dict[str, Any]:
    """단일 시트 celldata 추출."""
    if sheet_name not in EDITOR_BOUNDS:
        raise ValueError(f"지원하지 않는 시트: {sheet_name}")
    r1, r2, c1, c2 = EDITOR_BOUNDS[sheet_name]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"시트 없음: {sheet_name}")
        ws = wb[sheet_name]
        celldata: list[dict[str, Any]] = []
        for row in ws.iter_rows(
            min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=False
        ):
            for cell in row:
                if not hasattr(cell, "row"):
                    continue
                vobj = _fortune_cell_from_openpyxl(cell)
                if not vobj:
                    continue
                celldata.append(
                    {
                        "r": int(cell.row) - 1,
                        "c": int(cell.column) - 1,
                        "v": vobj,
                    }
                )
        return {
            "name": sheet_name,
            "celldata": celldata,
            "row": r2,
            "column": c2,
            "status": 1 if active else 0,
        }
    finally:
        wb.close()


def get_editor_sheet_cached(
    upload_dir: Path, batch_id: int, xlsx_path: Path, sheet_name: str, *, active: bool = False
) -> dict[str, Any]:
    cache = editor_cache_path(upload_dir, batch_id, sheet_name)
    if _cache_is_fresh(cache, xlsx_path):
        hit = load_cached_sheet(cache)
        if hit:
            hit["status"] = 1 if active else 0
            return hit
    sheet = export_editor_sheet(xlsx_path, sheet_name, active=active)
    write_cached_sheet(cache, sheet)
    return sheet


def export_editor_sheets(path: Path) -> dict[str, Any]:
    """현황판·2026 전체 (레거시)."""
    sheets_out = []
    for i, name in enumerate(EDITOR_SHEET_NAMES):
        sheets_out.append(export_editor_sheet(path, name, active=(i == 0)))
    return {"sheets": sheets_out}


def _apply_vobj_to_cell(cell: Any, vobj: dict[str, Any]) -> None:
    if not isinstance(vobj, dict):
        return
    if vobj.get("f"):
        cell.value = str(vobj["f"])
        return
    if "v" in vobj:
        cell.value = vobj["v"]


def apply_editor_sheets(path: Path, sheets_payload: list[dict[str, Any]]) -> int:
    """편집된 celldata를 원본 xlsx 해당 시트에 반영. 반환: 갱신 셀 수."""
    by_name = {s.get("name"): s for s in sheets_payload if isinstance(s, dict) and s.get("name")}
    if not by_name:
        return 0

    wb = openpyxl.load_workbook(path, data_only=False)
    updated = 0
    try:
        for name, bounds in EDITOR_BOUNDS.items():
            if name not in by_name or name not in wb.sheetnames:
                continue
            ws = wb[name]
            r1, r2, c1, c2 = bounds
            touched: set[tuple[int, int]] = set()
            for item in by_name[name].get("celldata") or []:
                if not isinstance(item, dict):
                    continue
                r = int(item.get("r", -1))
                c = int(item.get("c", -1))
                er, ec = r + 1, c + 1
                if er < r1 or er > r2 or ec < c1 or ec > c2:
                    continue
                vobj = item.get("v")
                if not isinstance(vobj, dict):
                    continue
                cell = ws.cell(er, ec)
                _apply_vobj_to_cell(cell, vobj)
                touched.add((er, ec))
            updated += len(touched)
        wb.save(path)
    finally:
        wb.close()
    return updated


def invalidate_editor_cache(upload_dir: Path, batch_id: int) -> None:
    for name in EDITOR_SHEET_NAMES:
        p = editor_cache_path(upload_dir, batch_id, name)
        try:
            p.unlink(missing_ok=True)
        except OSError:
            pass
